import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE & DESIGN SYSTEM
# ==============================================================================
C_NAVY_DARK = RGBColor(15, 23, 42)       # #0F172A (Title slide bg, dark text)
C_NAVY_MED = RGBColor(30, 41, 59)        # #1E293B (Card titles, primary headers)
C_SLATE_MUTED = RGBColor(100, 116, 139)  # #64748B (Subtitles, labels)
C_CARD_BG = RGBColor(248, 250, 252)      # #F8FAFC (Card background)
C_CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0 (Card border)
C_WHITE = RGBColor(255, 255, 255)
C_TEAL = RGBColor(13, 148, 136)          # #0D9488 (Solar, success, accent)
C_TEAL_LIGHT = RGBColor(240, 253, 250)   # #F0FDFA (Badge bg)
C_BLUE = RGBColor(37, 99, 235)           # #2563EB (Algorithms, routing)
C_AMBER = RGBColor(217, 119, 6)          # #D97706 (Warning, draining)
C_RED = RGBColor(220, 38, 38)            # #DC2626 (Failure, trap)
C_CODE_BG = RGBColor(241, 245, 249)      # #F1F5F9 (Code block bg)

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"
FONT_CODE = "Consolas"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(BASE_DIR, "results")
EXPLAINER_DIR = os.path.join(BASE_DIR, "report", "explainer_images")
EQ_DIR = os.path.join(BASE_DIR, "rep2_eq_pngs")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title, category, slide_num, total_slides=21):
    """Draws standardized corporate header, category pill badge, and running footer."""
    # Top Category Pill Badge
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.2), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_TEAL_LIGHT
    pill.line.color.rgb = C_TEAL
    pill.line.width = Pt(0.75)
    tf = pill.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
    tf_t = tb.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK

    # Top separator line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = C_CARD_BORDER
    line.line.fill.background()

    # Running Footer on slide
    f_tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(10.0), Inches(0.35))
    f_tf = f_tb.text_frame
    f_p = f_tf.paragraphs[0]
    f_p.text = "Adaptive Routing in EH-WSNs  |  Santhosh G (25BEL0016), Thilang T V (25BEL0036), Ananthakrishnan R (25BEL0005)"
    f_p.font.name = FONT_BODY
    f_p.font.size = Pt(8.5)
    f_p.font.color.rgb = C_SLATE_MUTED

    # Slide Number
    num_tb = slide.shapes.add_textbox(Inches(11.3), Inches(7.0), Inches(1.2), Inches(0.35))
    num_tf = num_tb.text_frame
    num_p = num_tf.paragraphs[0]
    num_p.text = f"{slide_num} / {total_slides}"
    num_p.font.name = FONT_BODY
    num_p.font.size = Pt(8.5)
    num_p.font.bold = True
    num_p.font.color.rgb = C_SLATE_MUTED
    num_p.alignment = PP_ALIGN.RIGHT

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
    """Draws a rounded card container for content grouping."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

def add_speaker_notes(slide, notes_text):
    """Attaches complete speaker notes to the slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text.strip()

print("Presentation template framework initialized.")
