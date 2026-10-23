import os
import sys
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE & DESIGN SYSTEM (Professional, Clean, Modern Academic)
# ==============================================================================
C_NAVY_DARK = RGBColor(15, 23, 42)       # #0F172A (Primary Dark)
C_NAVY_MED = RGBColor(30, 41, 59)        # #1E293B (Card headers)
C_SLATE_MUTED = RGBColor(100, 116, 139)  # #64748B (Secondary / Muted text)
C_CARD_BG = RGBColor(248, 250, 252)      # #F8FAFC (Card background)
C_CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0 (Subtle border)
C_WHITE = RGBColor(255, 255, 255)
C_TEAL = RGBColor(13, 148, 136)          # #0D9488 (Solar, Success, Accent)
C_TEAL_LIGHT = RGBColor(240, 253, 250)   # #F0FDFA (Teal tint bg)
C_BLUE = RGBColor(37, 99, 235)           # #2563EB (Tech / Algorithm)
C_BLUE_LIGHT = RGBColor(239, 246, 255)   # #EFF6FF (Blue tint bg)
C_AMBER = RGBColor(217, 119, 6)          # #D97706 (Warning, Draining)
C_AMBER_LIGHT = RGBColor(254, 243, 199)  # #FEF3C7 (Amber tint bg)
C_RED = RGBColor(220, 38, 38)            # #DC2626 (Failure, Brownout)
C_CODE_BG = RGBColor(241, 245, 249)      # #F1F5F9 (Code block bg)
C_DARK_TERMINAL = RGBColor(17, 24, 39)   # #111827 (Terminal black)

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"
FONT_CODE = "Consolas"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(BASE_DIR, "results")
EXPLAINER_DIR = os.path.join(BASE_DIR, "report", "explainer_images")
EQ_DIR = os.path.join(BASE_DIR, "rep2_eq_pngs")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title, category, slide_num, total_slides=21):
    """Draws standardized humanized corporate header, category pill badge, and running footer."""
    # Top Category Pill Badge
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.6), Inches(0.3))
    pill.fill.solid()
    pill.fill.fore_color.rgb = C_TEAL_LIGHT
    pill.line.color.rgb = C_TEAL
    pill.line.width = Pt(0.75)
    tf = pill.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = C_TEAL
    p.alignment = PP_ALIGN.CENTER

    # Main Title (Conversational, punchy, clear)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.65))
    tf_t = tb.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = C_NAVY_DARK

    # Top separator line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.36), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = C_CARD_BORDER
    line.line.fill.background()

    # Running Footer on slide
    f_tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(10.2), Inches(0.35))
    f_tf = f_tb.text_frame
    f_p = f_tf.paragraphs[0]
    f_p.text = "Adaptive Routing in EH-WSNs  |  Santhosh G (25BEL0016), Thilang T V (25BEL0036), Ananthakrishnan R (25BEL0005)"
    f_p.font.name = FONT_BODY
    f_p.font.size = Pt(8.5)
    f_p.font.color.rgb = C_SLATE_MUTED

    # Slide Number
    num_tb = slide.shapes.add_textbox(Inches(11.3), Inches(7.0), Inches(1.2), Inches(0.35))
    num_tf = num_tb.text_frame
    num_p = num_tf.paragraphs[0]
    num_p.text = f"{slide_num} / {total_slides}"
    num_p.font.name = FONT_BODY
    num_p.font.size = Pt(8.5)
    num_p.font.bold = True
    num_p.font.color.rgb = C_SLATE_MUTED
    num_p.alignment = PP_ALIGN.RIGHT

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
    """Draws a rounded card container for content grouping."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

def add_speaker_notes(slide, notes_text):
    """Attaches complete, natural, humanized speaker notes to the slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text.strip()

def add_bullet(tf, bold_prefix, text, pt_size=11, space_after=6, text_color=C_NAVY_DARK):
    """Adds a formatted bullet point with bold lead-in."""
    p = tf.add_paragraph()
    p.space_after = Pt(space_after)
    if bold_prefix:
        r_bold = p.add_run()
        r_bold.text = bold_prefix + " "
        r_bold.font.name = FONT_BODY
        r_bold.font.size = Pt(pt_size)
        r_bold.font.bold = True
        r_bold.font.color.rgb = text_color
    r_text = p.add_run()
    r_text.text = text
    r_text.font.name = FONT_BODY
    r_text.font.size = Pt(pt_size)
    r_text.font.color.rgb = text_color

print("Humanized PPT builder ready.")
